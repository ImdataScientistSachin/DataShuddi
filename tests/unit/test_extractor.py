"""
DataShuddi -- Unit Tests: Node 1 Extractor.

Verifies text extraction from PDF, DOCX, and TXT files and handling of unsupported types.
Follows the AAA (Arrange, Act, Assert) testing pattern.
"""
from pathlib import Path

import pytest

from src.agent.nodes.extractor import extractor_node
from src.agent.state import DataShuddiState


@pytest.fixture
def txt_sample_file(tmp_path: Path) -> Path:
    """Fixture creating a temporary test TXT file."""
    f = tmp_path / "sample.txt"
    f.write_text("Line 1: Acme Annual Report\nLine 2: Section A\nLine 3: Financial Summary", encoding="utf-8")
    return f


@pytest.fixture
def docx_sample_file(tmp_path: Path) -> Path:
    """Fixture creating a temporary test DOCX file."""
    from docx import Document
    doc = Document()
    doc.add_paragraph("Acme DOCX Executive Summary Report.")
    doc.add_paragraph("Quarterly highlights show positive operating income.")
    f = tmp_path / "sample.docx"
    doc.save(str(f))
    return f


@pytest.fixture
def base_state(txt_sample_file: Path) -> DataShuddiState:
    """Fixture providing baseline DataShuddiState with a TXT file."""
    return {
        "doc_id": "hash_txt_123",
        "doc_path": str(txt_sample_file),
        "source_system": "local",
        "mime_type": "text/plain",
        "raw_text": "",
        "_pages": [],
        "cleaned_text": "",
        "masked_text": "",
        "header_footer_candidates": [],
        "pii_entities": [],
        "pii_count": 0,
        "validation_errors": [],
        "validation_passed": False,
        "retry_count": 0,
        "reflection_text": "",
        "reflector_confidence": 0.0,
        "needs_human_review": False,
        "hf_threshold": 0.30,
        "is_sanitized": True,
        "security_flags": [],
        "quality_score": 0.0,
        "node_timestamps": {},
        "processing_time_ms": 0,
        "_start_time": 0.0,
    }


class TestExtractorNode:
    """Unit tests for Node 1: Extractor."""

    def test_txt_extraction_returns_text(self, base_state: DataShuddiState) -> None:
        """Arrange: Valid TXT file. Act: Run extractor node. Assert: Text extracted into raw_text & _pages."""
        # Arrange
        state = base_state

        # Act
        result = extractor_node(state)

        # Assert
        assert len(result["raw_text"]) > 0
        assert len(result["_pages"]) == 1
        assert "Acme Annual Report" in result["raw_text"]
        assert "extractor" in result["node_timestamps"]

    def test_docx_extraction_returns_text(self, base_state: DataShuddiState, docx_sample_file: Path) -> None:
        """Arrange: Valid DOCX file. Act: Run extractor node. Assert: Text extracted into raw_text."""
        # Arrange
        state = base_state
        state["doc_path"] = str(docx_sample_file)
        state["mime_type"] = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        state["raw_text"] = ""
        state["_pages"] = []

        # Act
        result = extractor_node(state)

        # Assert
        assert len(result["raw_text"]) > 0
        assert len(result["_pages"]) == 1
        assert "Acme DOCX Executive Summary Report" in result["raw_text"]

    def test_unsupported_mime_returns_empty_raw_text(self, base_state: DataShuddiState) -> None:
        """Arrange: Unsupported MIME type. Act: Run node. Assert: Returns empty raw_text without crashing."""
        # Arrange
        state = base_state
        state["mime_type"] = "application/zip"

        # Act
        result = extractor_node(state)

        # Assert
        assert result["raw_text"] == ""
        assert result["_pages"] == []

    def test_nonexistent_file_caught_gracefully(self, base_state: DataShuddiState) -> None:
        """Arrange: Non-existent doc_path without inline text. Act: Run node. Assert: Returns empty text."""
        # Arrange
        state = base_state
        state["doc_path"] = "./non_existent_file_path_9999.pdf"
        state["mime_type"] = "application/pdf"
        state["raw_text"] = ""

        # Act
        result = extractor_node(state)

        # Assert
        assert result["raw_text"] == ""
        assert result["_pages"] == []

    def test_inline_raw_text_preserved_if_file_missing(self, base_state: DataShuddiState) -> None:
        """Arrange: Nonexistent doc_path but populated raw_text & _pages. Act: Run node. Assert: Preserved."""
        # Arrange
        state = base_state
        state["doc_path"] = "./missing_fixture_doc.txt"
        state["raw_text"] = "Preserved inline fixture text content."
        state["_pages"] = ["Preserved inline fixture text content."]

        # Act
        result = extractor_node(state)

        # Assert
        assert result["raw_text"] == "Preserved inline fixture text content."
        assert len(result["_pages"]) == 1
        assert "extractor" in result["node_timestamps"]

    def test_corrupted_pdf_bytes_handled_gracefully(self, base_state: DataShuddiState, tmp_path: Path) -> None:
        """Arrange: Invalid/corrupted PDF bytes on disk. Act: Run node. Assert: Exception caught gracefully."""
        # Arrange
        bad_pdf = tmp_path / "corrupted.pdf"
        bad_pdf.write_bytes(b"%PDF-1.4 Invalid corrupted PDF binary data stream")

        state = base_state
        state["doc_path"] = str(bad_pdf)
        state["mime_type"] = "application/pdf"
        state["raw_text"] = ""

        # Act
        result = extractor_node(state)

        # Assert
        assert result["raw_text"] == ""
        assert result["_pages"] == []
